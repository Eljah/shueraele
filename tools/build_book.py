"""Build the Russian mini-book from Markdown plus the assembly/source appendices.
The PDF uses a local LibreOffice installation. No network or font redistribution.
"""
from pathlib import Path
import argparse,re,subprocess,shutil,tempfile
from docx import Document
from docx.shared import Mm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
p=argparse.ArgumentParser();p.add_argument('--pdf',action='store_true');args=p.parse_args()
book=ROOT/'book';book.mkdir(exist_ok=True)
text=(book/'manuscript.md').read_text(encoding='utf8')
text+='\n\n## Приложение А. Паспорт нативной сборки\n\n'+(ROOT/'docs/ASSEMBLY_RU.md').read_text(encoding='utf8').replace('# Сборка исторического эффекта в Meta Spark\n','').replace('## ','### ')
text+='\n\n## Приложение Б. Источники\n\n'+(ROOT/'docs/SOURCES.md').read_text(encoding='utf8').replace('# Источники и границы проверки\n','')
(book/'mini-book.md').write_text(text,encoding='utf8')
doc=Document();sec=doc.sections[0];sec.page_width=Mm(148);sec.page_height=Mm(210)
sec.top_margin=Mm(17);sec.bottom_margin=Mm(17);sec.left_margin=Mm(16);sec.right_margin=Mm(14)
sec.header_distance=Mm(8);sec.footer_distance=Mm(8)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
 style=doc.styles[name];style.font.name='DejaVu Sans';style._element.rPr.rFonts.set(qn('w:eastAsia'),'DejaVu Sans')
normal=doc.styles['Normal'];normal.font.size=Pt(9.5);normal.paragraph_format.line_spacing=1.16;normal.paragraph_format.space_after=Pt(7)
for name,size in [('Title',25),('Heading 1',17),('Heading 2',12),('Heading 3',11)]:
 st=doc.styles[name];st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string('284A35');st.paragraph_format.keep_with_next=True
 if name=='Heading 1':st.paragraph_format.page_break_before=True
header=sec.header.paragraphs[0];header.text='НЕКРОЛАБОРАТОРИЯ  /  ШҮРӘЛЕ';header.style=doc.styles['Normal']
for r in header.runs:r.font.size=Pt(7);r.font.color.rgb=RGBColor.from_string('637861')
footer=sec.footer.paragraphs[0];footer.alignment=2
r=footer.add_run('Eljah/shueraele  ·  ');r.font.size=Pt(7)
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');footer._p.append(f)
code=False;buf=[];in_code=[]
def inline(par,text):
 for part in re.split(r'(\*\*.*?\*\*|`[^`]*`)',text):
  r=par.add_run(part[2:-2] if part.startswith('**') else part[1:-1] if part.startswith('`') else part)
  if part.startswith('**'):r.bold=True
  if part.startswith('`'):r.font.name='DejaVu Sans Mono';r.font.size=Pt(8)
def flush():
 if buf:inline(doc.add_paragraph(),' '.join(buf));buf.clear()
for line in text.splitlines():
 if line.startswith('```'):
  flush();code=not code
  if not code:
   p1=doc.add_paragraph();p1.paragraph_format.line_spacing=1;p1.paragraph_format.space_after=Pt(9)
   for i,s in enumerate(in_code):
    r=p1.add_run(('\n' if i else '')+s);r.font.name='DejaVu Sans Mono';r.font.size=Pt(6.7)
   shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EDF1E7');p1._p.get_or_add_pPr().append(shade);in_code.clear()
  continue
 if code:in_code.append(line);continue
 if not line.strip():flush();continue
 if line.startswith('!['):
  flush();m=re.match(r'!\[(.*?)\]\((.*?)\)',line)
  if m:
   path=(book/m.group(2)).resolve()
   if path.exists():doc.add_picture(str(path),width=Mm(116))
   else:doc.add_paragraph('[Иллюстрация ещё не собрана: '+m.group(2)+']')
   c=doc.add_paragraph(m.group(1));c.paragraph_format.space_after=Pt(10)
   for r in c.runs:r.italic=True;r.font.size=Pt(7.5)
  continue
 if line.startswith('# '):flush();doc.add_paragraph(line[2:],style='Title');continue
 if line.startswith('## '):
  flush();heading=line[3:];style='Subtitle' if heading.startswith('Шүрәле:') else 'Heading 1';doc.add_paragraph(heading,style=style);continue
 if line.startswith('### '):flush();doc.add_paragraph(line[4:],style='Heading 2');continue
 buf.append(line)
flush();doc.core_properties.title='Шүрәле. Платформа ушла. Рог остался.';doc.core_properties.author='Shurale Necrolab';doc.core_properties.subject='Исторический AR-комплект и проверяемый отладочный стенд'
out=book/'mini-book.docx';doc.save(out)
if args.pdf:
 exe=shutil.which('libreoffice') or shutil.which('soffice')
 if not exe:raise SystemExit('LibreOffice is required for --pdf; DOCX has been written.')
 with tempfile.TemporaryDirectory() as profile:
  subprocess.run([exe,'-env:UserInstallation='+Path(profile).as_uri(),'--headless','--convert-to','pdf','--outdir',str(book),str(out)],check=True,timeout=120)
 if not (book/'mini-book.pdf').exists():raise SystemExit('LibreOffice did not create the PDF.')
print(out)
